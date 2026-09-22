import io
from datetime import date, time

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import TestCase
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from employees.models import Department, Employee

from .models import ShiftPlan, ShiftPlanAssignment
from .services.roster_upload import apply_roster_upload, build_template_workbook


def make_upload(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Roster Upload"
    sheet.append(["Employee ID", "Name", "Department", "Current Plan", "Current Group", "Plan Name", "Group (A/B)", "Start Date (YYYY-MM-DD)"])
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    output.name = "upload.xlsx"
    return output


class RosterUploadServiceTests(TestCase):
    def setUp(self):
        self.dept = Department.objects.create(name="Extrusion")
        self.day = Shift = __import__("attendance.models", fromlist=["Shift"]).Shift
        self.day = Shift.objects.create(name="Day", start_time=time(7), end_time=time(19))
        self.night = Shift.objects.create(name="Night", start_time=time(19), end_time=time(7), is_overnight=True)
        self.rotation = ShiftPlan.objects.create(name="Extrusion Rotation", kind="rotation", day_shift=self.day, night_shift=self.night, anchor_monday=date(2026, 9, 21))
        self.perm_day = ShiftPlan.objects.create(name="Permanent Day", kind="fixed", shift=self.day, working_weekdays=[0, 1, 2, 3, 4, 5])
        self.alice = Employee.objects.create(employee_id="E001", first_name="Alice", last_name="A", department=self.dept)
        self.bob = Employee.objects.create(employee_id="E002", first_name="Bob", last_name="B", department=self.dept)

    def test_template_lists_every_active_employee_with_current_plan(self):
        ShiftPlanAssignment.objects.create(employee=self.alice, plan=self.rotation, group="A", start_date=date(2026, 9, 1))
        content = build_template_workbook()
        workbook = load_workbook(io.BytesIO(content))
        self.assertIn("Roster Upload", workbook.sheetnames)
        self.assertIn("Plans (reference)", workbook.sheetnames)
        rows = list(workbook["Roster Upload"].iter_rows(values_only=True))
        by_id = {row[0]: row for row in rows[1:]}
        self.assertEqual(set(by_id), {"E001", "E002"})
        self.assertEqual(by_id["E001"][3], "Extrusion Rotation")  # Current Plan
        self.assertEqual(by_id["E001"][4], "A")  # Current Group
        self.assertIsNone(by_id["E002"][3])  # Bob has no plan yet (blank cell)

    def test_dry_run_reports_without_changing_anything(self):
        upload = make_upload([["E001", "", "", "", "", "Permanent Day", "", ""]])
        report = apply_roster_upload(upload, dry_run=True)
        self.assertTrue(report.dry_run)
        self.assertEqual(report.changes, 1)
        self.assertFalse(report.issues)
        self.assertFalse(ShiftPlanAssignment.objects.filter(employee=self.alice).exists())

    def test_apply_assigns_a_fixed_plan(self):
        upload = make_upload([["E001", "", "", "", "", "Permanent Day", "", "2026-10-01"]])
        report = apply_roster_upload(upload, dry_run=False, actor="hr")
        self.assertEqual(report.changes, 1)
        assignment = ShiftPlanAssignment.objects.get(employee=self.alice)
        self.assertEqual(assignment.plan, self.perm_day)
        self.assertEqual(assignment.start_date, date(2026, 10, 1))
        self.assertEqual(assignment.assigned_by, "hr")

    def test_apply_assigns_a_rotation_plan_with_group(self):
        upload = make_upload([["E002", "", "", "", "", "Extrusion Rotation", "B", ""]])
        apply_roster_upload(upload, dry_run=False)
        assignment = ShiftPlanAssignment.objects.get(employee=self.bob)
        self.assertEqual(assignment.plan, self.rotation)
        self.assertEqual(assignment.group, "B")

    def test_rotation_plan_without_group_is_an_issue(self):
        upload = make_upload([["E002", "", "", "", "", "Extrusion Rotation", "", ""]])
        report = apply_roster_upload(upload, dry_run=True)
        self.assertEqual(report.changes, 0)
        self.assertEqual(len(report.issues), 1)
        self.assertIn("Group A or B", report.issues[0].reason)

    def test_blank_plan_name_leaves_the_row_unchanged(self):
        upload = make_upload([["E001", "", "", "", "", "", "", ""]])
        report = apply_roster_upload(upload, dry_run=True)
        self.assertEqual(report.changes, 0)
        self.assertEqual(report.unchanged, 1)
        self.assertFalse(report.issues)

    def test_unknown_employee_is_an_issue(self):
        upload = make_upload([["NOBODY", "", "", "", "", "Permanent Day", "", ""]])
        report = apply_roster_upload(upload, dry_run=True)
        self.assertEqual(len(report.issues), 1)
        self.assertIn("No active employee", report.issues[0].reason)

    def test_unknown_plan_name_is_an_issue(self):
        upload = make_upload([["E001", "", "", "", "", "Made Up Plan", "", ""]])
        report = apply_roster_upload(upload, dry_run=True)
        self.assertEqual(len(report.issues), 1)
        self.assertIn("No active plan named", report.issues[0].reason)

    def test_two_people_on_the_same_plan_are_grouped_in_one_bucket(self):
        upload = make_upload([
            ["E001", "", "", "", "", "Permanent Day", "", "2026-10-01"],
            ["E002", "", "", "", "", "Permanent Day", "", "2026-10-01"],
        ])
        report = apply_roster_upload(upload, dry_run=False)
        self.assertEqual(report.changes, 2)
        self.assertEqual(len(report.by_plan), 1)
        self.assertEqual(ShiftPlanAssignment.objects.filter(plan=self.perm_day).count(), 2)


class RosterUploadAPITests(TestCase):
    def setUp(self):
        self.dept = Department.objects.create(name="Extrusion")
        from .models import Shift

        self.day = Shift.objects.create(name="Day", start_time=time(7), end_time=time(19))
        self.perm_day = ShiftPlan.objects.create(name="Permanent Day", kind="fixed", shift=self.day, working_weekdays=[0, 1, 2, 3, 4, 5])
        self.alice = Employee.objects.create(employee_id="E001", first_name="Alice", last_name="A", department=self.dept)

        user_model = get_user_model()
        self.user = user_model.objects.create_user(username="hr", password="pw")
        self.user.user_permissions.add(Permission.objects.get(codename="manage_shifts"))
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def test_template_download(self):
        response = self.client.get("/api/attendance/shift-roster/template/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertIn("E001", [row[0] for row in workbook["Roster Upload"].iter_rows(values_only=True, min_row=2)])

    def test_upload_dry_run_then_apply(self):
        upload = make_upload([["E001", "", "", "", "", "Permanent Day", "", ""]])
        response = self.client.post("/api/attendance/shift-roster/upload/", {"file": upload}, format="multipart")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["changes"], 1)
        self.assertTrue(response.data["dry_run"])
        self.assertFalse(ShiftPlanAssignment.objects.exists())

        upload.seek(0)
        response = self.client.post("/api/attendance/shift-roster/upload/", {"file": upload, "dry_run": "false"}, format="multipart")
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.data["dry_run"])
        self.assertTrue(ShiftPlanAssignment.objects.filter(employee=self.alice, plan=self.perm_day).exists())

    def test_upload_requires_manage_shifts_permission(self):
        self.user.user_permissions.clear()
        upload = make_upload([["E001", "", "", "", "", "Permanent Day", "", ""]])
        response = self.client.post("/api/attendance/shift-roster/upload/", {"file": upload}, format="multipart")
        self.assertEqual(response.status_code, 403)
