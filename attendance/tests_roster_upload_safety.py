import io

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from attendance.services.roster_upload import build_template_workbook
from employees.models import Employee


class RosterTemplateFormulaSafetyTests(TestCase):
    """2026-09-25 review, S-07: a staff number or name starting with = must not become a formula in the export."""

    def test_untrusted_strings_stay_text(self):
        Employee.objects.create(employee_id="=1+1", first_name="=cmd|' /C calc'!A0", last_name="Person", status="active")
        Employee.objects.create(employee_id="000777", first_name="+Plus", last_name="@Name", status="active")
        sheet = load_workbook(io.BytesIO(build_template_workbook())).active
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                self.assertNotEqual(cell.data_type, "f", f"{cell.coordinate} is a formula: {cell.value!r}")
        self.assertIn("=1+1", [row[0].value for row in sheet.iter_rows(min_row=2)])


class StaffNumberFormulaValidationTests(APITestCase):
    def test_a_staff_number_that_starts_like_a_formula_is_rejected(self):
        user = get_user_model().objects.create_user(username="editor", password="x")
        user.user_permissions.add(Permission.objects.get(codename="add_employee", content_type__app_label="employees"))
        self.client.force_authenticate(user)
        for bad in ("=1+1", "@SUM(A1)", "+44", "-1"):
            response = self.client.post("/api/employees/", {"employee_id": bad, "first_name": "Ada", "last_name": "Okafor"}, format="json")
            self.assertEqual(response.status_code, 400, bad)
        self.assertFalse(Employee.objects.exists())
