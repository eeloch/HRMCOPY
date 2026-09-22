"""Excel upload for assigning shift plans in bulk, as an alternative to the "assign a department" screen.

Download the template and it already lists every active employee with their staff number, name, department
and whatever plan/group they are on today. Fill in Plan Name (and Group, for a rotation plan) only for the
people who should change, leave the rest blank, and upload it back. Nobody is touched unless a row names a
plan for them.
"""

import io
from dataclasses import dataclass, field
from datetime import date as date_cls
from datetime import datetime as datetime_cls

from django.db import transaction
from django.db.models import Q
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from attendance.models import ShiftPlan, ShiftPlanAssignment
from attendance.services.shift_plans import assign_plan
from employees.models import Employee

SHEET_TITLE = "Roster Upload"
HEADERS = [
    "Employee ID",
    "Name",
    "Department",
    "Current Plan",
    "Current Group",
    "Plan Name",
    "Group (A/B)",
    "Start Date (YYYY-MM-DD)",
]


def _current_assignments(today=None):
    """The one active ShiftPlanAssignment per employee, as of today."""
    today = today or timezone.localdate()
    rows = (
        ShiftPlanAssignment.objects.select_related("plan")
        .filter(start_date__lte=today)
        .filter(Q(end_date__isnull=True) | Q(end_date__gte=today))
        .order_by("employee_id", "-start_date")
    )
    result = {}
    for row in rows:
        result.setdefault(row.employee_id, row)
    return result


def build_template_workbook():
    """One row per active employee, already carrying their details and current plan, so HR only needs to
    type the Plan Name (and Group, for a rotation plan) for anyone whose roster should change."""
    employees = Employee.objects.filter(status="active").select_related("department").order_by("employee_id")
    current = _current_assignments()
    plans = list(ShiftPlan.objects.filter(active=True).order_by("name"))
    plan_names = [plan.name for plan in plans]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_TITLE
    sheet.append(HEADERS)
    row_count = 0
    for employee in employees:
        row_count += 1
        assignment = current.get(employee.pk)
        sheet.append([
            employee.employee_id,
            employee.full_name,
            employee.department.name if employee.department else "",
            assignment.plan.name if assignment else "",
            assignment.group if assignment else "",
            "",
            "",
            "",
        ])
    for column, width in zip("ABCDEFGH", (12, 26, 20, 24, 12, 24, 14, 24)):
        sheet.column_dimensions[column].width = width

    reference = workbook.create_sheet("Plans (reference)")
    reference.append(["Plan Name", "Kind"])
    for plan in plans:
        reference.append([plan.name, plan.get_kind_display()])
    reference.column_dimensions["A"].width = 26
    reference.column_dimensions["B"].width = 22

    last_row = row_count + 1
    if plan_names:
        joined = ",".join(plan_names)
        formula = f'"{joined}"' if len(joined) <= 250 else f"'Plans (reference)'!$A$2:$A${len(plan_names) + 1}"
        plan_validation = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True)
        plan_validation.error = 'Pick a plan name from the "Plans (reference)" sheet.'
        sheet.add_data_validation(plan_validation)
        plan_validation.add(f"F2:F{last_row}")

    group_validation = DataValidation(type="list", formula1='"A,B"', allow_blank=True, showErrorMessage=True)
    group_validation.error = "Group must be A or B (only for a rotation plan)."
    sheet.add_data_validation(group_validation)
    group_validation.add(f"G2:G{last_row}")

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@dataclass
class RosterUploadIssue:
    row: int
    employee_id: str
    reason: str


@dataclass
class RosterUploadReport:
    dry_run: bool = True
    rows_in_file: int = 0
    changes: int = 0
    unchanged: int = 0
    by_plan: dict = field(default_factory=dict)
    issues: list = field(default_factory=list)

    def as_dict(self):
        data = {"dry_run": self.dry_run, "rows_in_file": self.rows_in_file, "changes": self.changes, "unchanged": self.unchanged, "by_plan": self.by_plan}
        data["issues"] = [issue.__dict__ for issue in self.issues]
        return data


def _parse_start_date(value):
    if value in (None, ""):
        return timezone.localdate()
    if isinstance(value, datetime_cls):
        return value.date()
    if isinstance(value, date_cls):
        return value
    try:
        return date_cls.fromisoformat(str(value).strip())
    except ValueError:
        raise ValueError(f'"{value}" is not a date (use YYYY-MM-DD).')


def _read_rows(file):
    workbook = load_workbook(file, data_only=True)
    sheet = workbook[SHEET_TITLE] if SHEET_TITLE in workbook.sheetnames else workbook.worksheets[0]
    all_rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    if not all_rows:
        raise ValueError("The sheet is empty.")
    header = [str(value or "").strip().lower() for value in all_rows[0]]

    def column(name):
        try:
            return header.index(name)
        except ValueError:
            return None

    idx = {
        "employee_id": column("employee id"),
        "plan": column("plan name"),
        "group": column("group (a/b)"),
        "start": column("start date (yyyy-mm-dd)"),
    }
    if idx["employee_id"] is None or idx["plan"] is None:
        raise ValueError('The sheet needs "Employee ID" and "Plan Name" columns.')
    return all_rows[1:], idx


def apply_roster_upload(file, *, dry_run=True, actor=""):
    """Read the filled-in template and either report what it would change (dry_run) or apply it."""
    rows, idx = _read_rows(file)
    employees_by_number = {employee.employee_id: employee for employee in Employee.objects.filter(status="active")}
    plans_by_name = {plan.name.strip().lower(): plan for plan in ShiftPlan.objects.filter(active=True)}
    report = RosterUploadReport(dry_run=dry_run, rows_in_file=len(rows))

    # Group people by (plan, group, start date) so each combination is applied in one call.
    buckets = {}
    for position, row in enumerate(rows, start=2):
        get = lambda key: (row[idx[key]] if idx[key] is not None and idx[key] < len(row) else None)
        employee_id = str(get("employee_id") or "").strip()
        plan_name = str(get("plan") or "").strip()
        if not employee_id and not plan_name:
            continue
        if not employee_id:
            report.issues.append(RosterUploadIssue(position, "", "No employee ID."))
            continue
        employee = employees_by_number.get(employee_id)
        if employee is None:
            report.issues.append(RosterUploadIssue(position, employee_id, "No active employee with this staff number."))
            continue
        if not plan_name:
            report.unchanged += 1
            continue  # nothing chosen for this person: leave their roster as it is
        plan = plans_by_name.get(plan_name.strip().lower())
        if plan is None:
            report.issues.append(RosterUploadIssue(position, employee_id, f'No active plan named "{plan_name}".'))
            continue
        group = str(get("group") or "").strip().upper()
        if plan.kind == "rotation" and group not in ("A", "B"):
            report.issues.append(RosterUploadIssue(position, employee_id, "A rotation plan needs Group A or B."))
            continue
        if plan.kind == "fixed":
            group = ""
        try:
            start = _parse_start_date(get("start"))
        except ValueError as error:
            report.issues.append(RosterUploadIssue(position, employee_id, str(error)))
            continue
        buckets.setdefault((plan.pk, group, start), []).append(employee)

    def describe(key):
        plan_id, group, start = key
        plan = plans_by_name_by_id.get(plan_id)
        label = plan.name + (f" (Group {group})" if group else "")
        return f"{label} from {start.isoformat()}"

    plans_by_name_by_id = {plan.pk: plan for plan in plans_by_name.values()}
    report.changes = sum(len(people) for people in buckets.values())
    report.by_plan = {describe(key): len(people) for key, people in buckets.items()}

    if not dry_run:
        with transaction.atomic():
            for (plan_id, group, start), people in buckets.items():
                assign_plan(people, plans_by_name_by_id[plan_id], group=group, start_date=start, actor=actor)

    return report
