import re
from collections import defaultdict
from dataclasses import dataclass, field

from django.db import transaction
from openpyxl import load_workbook

from audit.models import AuditSeverity
from audit.services import AuditService
from employees.models import Employee

from .models import Building, BuildingKind, Room, RoomAssignment

BED_PATTERN = re.compile(r"^(?P<room>.*?)\s*-\s*Bed\s*(?P<bed>\d+)\s*$", re.IGNORECASE)
MAIN_HOSTEL = "Main Hostel"
BUILDING_ORDER = [MAIN_HOSTEL, "Lodge 1", "Lodge 2", "Security Quarters", "Admin"]


def split_room_name(text):
    """'Lodge 2 Rm005' -> ('Lodge 2', 'Rm005'). Buildings come from how rooms are named:
    Lodge 1/2, Security, Admin, and plain 'Room 301' (also written 'Room301') is the main hostel."""
    room = re.sub(r"\s+", " ", str(text or "")).strip()
    lodge = re.match(r"^(Lodge\s*\d+)\s*(.+)$", room, re.IGNORECASE)
    if lodge:
        return re.sub(r"\s+", " ", lodge.group(1)).title(), lodge.group(2).strip()
    security = re.match(r"^Security\s*(.+)$", room, re.IGNORECASE)
    if security:
        return "Security Quarters", security.group(1).strip()
    admin = re.match(r"^Admin\s*(.+)$", room, re.IGNORECASE)
    if admin:
        return "Admin", admin.group(1).strip()
    number = re.match(r"^Room\s*(\d+)$", room, re.IGNORECASE)
    if number:
        return MAIN_HOSTEL, f"Room {number.group(1)}"
    return MAIN_HOSTEL, room


def parse_room_label(label):
    """'Room 301 - Bed 1' -> ('Main Hostel', 'Room 301', 1)."""
    match = BED_PATTERN.match(str(label or "").strip())
    if not match:
        return None
    building, room = split_room_name(match.group("room"))
    return building, room, int(match.group("bed"))


def gender_of(value):
    text = str(value or "").strip().lower()
    return "male" if text in ("male", "m") else "female" if text in ("female", "f") else ""


def room_text(room):
    """The short text kept on the employee record: the room without 'Main Hostel'."""
    return room.name if room.building.name == MAIN_HOSTEL else f"{room.building.name} {room.name}"


def sync_employee_flags(employee):
    """Keep the older 'lives in hostel / outside' fields on the employee in step with where they are placed."""
    assignment = getattr(employee, "room_assignment", None)
    inside, outside, text, address = False, False, "", ""
    if assignment is not None:
        building = assignment.room.building
        if building.kind == BuildingKind.COMPANY:
            inside, text = True, room_text(assignment.room)
        else:
            outside, address = True, building.address or building.name
    else:
        inside, outside, text, address = employee.lives_in_company_hostel, employee.lives_in_external_accommodation, employee.hostel_room_number, employee.external_accommodation_address
    Employee.objects.filter(pk=employee.pk).update(lives_in_company_hostel=inside, lives_in_external_accommodation=outside, hostel_room_number=text if inside else "", external_accommodation_address=address if outside else "")


class AccommodationService:
    @staticmethod
    def assign(employee, room, *, bed_number=None, actor=None, enforce_capacity=True):
        with transaction.atomic():
            if employee.status != "active":
                raise ValueError("Only active staff can be given a room.")
            room = Room.objects.select_for_update().select_related("building").get(pk=room.pk)
            if not room.active:
                raise ValueError("This room is not in use.")
            if room.gender and employee.gender and room.gender != employee.gender:
                raise ValueError(f"{room.name} is a {room.gender} room and {employee.full_name} is {employee.gender}.")
            if bed_number is not None:
                taken = RoomAssignment.objects.filter(room=room, bed_number=bed_number).exclude(employee=employee).select_related("employee").first()
                if taken:
                    raise ValueError(f"Bed {bed_number} in {room.name} is already taken by {taken.employee.full_name}.")
            elif enforce_capacity and room.capacity is not None and RoomAssignment.objects.filter(room=room).exclude(employee=employee).count() >= room.capacity:
                raise ValueError(f"{room.name} is full ({room.capacity} of {room.capacity}).")
            existing = RoomAssignment.objects.filter(employee=employee).select_related("room__building").first()
            previous = f"{existing.room.building.name} {existing.room.name}" if existing else None
            assignment, _ = RoomAssignment.objects.update_or_create(employee=employee, defaults={"room": room, "bed_number": bed_number, "assigned_by": actor})
            sync_employee_flags(Employee.objects.get(pk=employee.pk))
            AuditService.log(event_type="accommodation.assigned", module="accommodation", employee=employee, actor=actor, object=assignment, severity=AuditSeverity.INFO, title="Room assigned", description=f"{employee.full_name} placed in {room.building.name} {room.name}" + (f", bed {bed_number}" if bed_number else "") + (f" (was {previous})." if previous else "."), metadata={"room_id": room.pk, "bed": bed_number})
        return assignment

    @staticmethod
    def unassign(employee, *, actor=None, outside=False):
        """Take someone out of their room. `outside` records that they live in accommodation outside the company."""
        with transaction.atomic():
            existing = RoomAssignment.objects.filter(employee=employee).select_related("room__building").first()
            if existing:
                label = f"{existing.room.building.name} {existing.room.name}"
                existing.delete()
                AuditService.log(event_type="accommodation.unassigned", module="accommodation", employee=employee, actor=actor, severity=AuditSeverity.INFO, title="Room vacated", description=f"{employee.full_name} left {label}.")
            employee = Employee.objects.get(pk=employee.pk)
            Employee.objects.filter(pk=employee.pk).update(lives_in_company_hostel=False, hostel_room_number="", lives_in_external_accommodation=outside, external_accommodation_address="" if not outside else employee.external_accommodation_address)

    @staticmethod
    def mark_inside_without_room(employee, written_room=""):
        """In company accommodation, but the room they were given is not one of our rooms (yet)."""
        RoomAssignment.objects.filter(employee=employee).delete()
        Employee.objects.filter(pk=employee.pk).update(lives_in_company_hostel=True, lives_in_external_accommodation=False, hostel_room_number=str(written_room or "")[:50], external_accommodation_address="")


def find_room(building_name, room_name):
    return Room.objects.select_related("building").filter(building__name=building_name, name=room_name).first()


def apply_import_placement(employee, placement, room_label, *, actor=None):
    """Place one employee as the Bulk Import's Accommodation / ROOM ALLOCATED columns say. Rooms are never
    created here: they come from the accommodation tracker or the Add Room button.
    Returns inside, inside_no_bed, unknown_room, gender_mismatch, vacated, outside or none (None = left alone)."""
    if not placement:
        return None
    parsed = parse_room_label(room_label) if room_label else None
    if placement == "inside" and parsed:
        building_name, room_name, bed = parsed
        room = find_room(building_name, room_name)
        if employee.status != "active":
            AccommodationService.unassign(employee, actor=actor)  # they have left: the bed is free again
            return "vacated"
        if room is None:
            AccommodationService.mark_inside_without_room(employee, room_label)
            return "unknown_room"
        try:
            AccommodationService.assign(employee, room, bed_number=bed, actor=actor, enforce_capacity=False)
            return "inside"
        except ValueError as error:
            if "is a " in str(error) and " room and " in str(error):
                AccommodationService.mark_inside_without_room(employee, room_label)
                return "gender_mismatch"
            AccommodationService.assign(employee, room, actor=actor, enforce_capacity=False)  # bed already taken: place without a bed number
            return "inside_no_bed"
    AccommodationService.unassign(employee, actor=actor, outside=placement == "outside" and employee.status == "active")
    return "outside" if placement == "outside" else "none"


# ---- reading spreadsheets ----------------------------------------------------------------

def _sheet_rows(file, sheet_hint, needed_headers=None):
    """All rows of the sheet whose name contains sheet_hint (.xlsx via openpyxl, .xlsb via pyxlsb)."""
    name = str(getattr(file, "name", "") or "").lower()
    if name.endswith(".xlsb"):
        from pyxlsb import open_workbook

        book = open_workbook(file)
        for sheet_name in book.sheets:
            if sheet_hint in sheet_name.lower():
                with book.get_sheet(sheet_name) as sheet:
                    return [[cell.v for cell in row] for row in sheet.rows()]
        raise ValueError(f'No sheet named "{sheet_hint}" was found in that file.')
    workbook = load_workbook(file, data_only=True)
    for sheet in workbook.worksheets:
        if sheet_hint in sheet.title.lower():
            return [list(row) for row in sheet.iter_rows(values_only=True)]
    raise ValueError(f'No sheet named "{sheet_hint}" was found in that file.')


def read_room_setup(file):
    """[(building, room, gender, beds)] from the tracker's Room Setup sheet: a MALE table and a FEMALE
    table side by side, each with Room No and Beds (capacity)."""
    rows = _sheet_rows(file, "room setup")
    header_index = next((i for i, row in enumerate(rows) if sum(1 for v in row if str(v or "").strip().lower() == "room no") >= 1), None)
    if header_index is None:
        raise ValueError('The Room Setup sheet has no "Room No" column.')
    header = rows[header_index]
    room_cols = [i for i, v in enumerate(header) if str(v or "").strip().lower() == "room no"]
    title_row = rows[header_index - 1] if header_index else []
    result, seen = [], set()
    for position, col in enumerate(room_cols):
        # The table's title ("MALE ROOMS ..." / "FEMALE ROOMS ...") says which gender it is.
        title = " ".join(str(title_row[c]) for c in range(col, min(col + 4, len(title_row))) if c < len(title_row) and title_row[c]).lower()
        gender = "female" if "female" in title else "male" if "male" in title else ("male", "female")[min(position, 1)]
        cap_col = next((c for c in range(col + 1, min(col + 4, len(header))) if "bed" in str(header[c] or "").lower()), col + 1)
        for row in rows[header_index + 1:]:
            if col >= len(row) or not str(row[col] or "").strip():
                continue
            text = str(row[col]).strip()
            if text.lower().startswith("total"):
                continue
            beds = row[cap_col] if cap_col < len(row) else None
            if not isinstance(beds, (int, float)) or beds <= 0:
                continue
            building, name = split_room_name(text)
            if (building, name) in seen:
                continue
            seen.add((building, name))
            result.append((building, name, gender, int(beds)))
    if not result:
        raise ValueError("No rooms were found in the Room Setup sheet.")
    return result


@dataclass
class RoomSetupReport:
    dry_run: bool = True
    rooms_in_file: int = 0
    male_rooms: int = 0
    male_beds: int = 0
    female_rooms: int = 0
    female_beds: int = 0
    created: list = field(default_factory=list)
    updated: list = field(default_factory=list)
    unchanged: int = 0
    removed: list = field(default_factory=list)
    kept_with_people: list = field(default_factory=list)
    gender_conflicts: list = field(default_factory=list)

    def as_dict(self):
        return self.__dict__.copy()


def import_room_setup(file, *, dry_run, actor=None):
    """Make the rooms match the accommodation tracker exactly: rooms in it are added or updated (gender and
    beds), and company rooms that are not in it are removed. A room that still has people in it is kept and
    reported instead, so nobody loses their place by accident."""
    wanted = read_room_setup(file)
    report = RoomSetupReport(dry_run=dry_run, rooms_in_file=len(wanted))
    for _b, _n, gender, beds in wanted:
        if gender == "male":
            report.male_rooms, report.male_beds = report.male_rooms + 1, report.male_beds + beds
        else:
            report.female_rooms, report.female_beds = report.female_rooms + 1, report.female_beds + beds
    with transaction.atomic():
        keep = set()
        for building_name, name, gender, beds in wanted:
            building, _ = Building.objects.get_or_create(name=building_name, defaults={"kind": BuildingKind.COMPANY, "sort_order": BUILDING_ORDER.index(building_name) if building_name in BUILDING_ORDER else 50})
            room = Room.objects.filter(building=building, name=name).first()
            label = f"{building_name} {name}"
            if room is None:
                room = Room.objects.create(building=building, name=name, gender=gender, capacity=beds, capacity_estimated=False)
                report.created.append(label)
            else:
                changes = []
                if room.gender != gender:
                    changes.append(f"gender {room.gender or 'not set'} to {gender}")
                    room.gender = gender
                if room.capacity != beds:
                    changes.append(f"beds {room.capacity if room.capacity is not None else 'not set'} to {beds}")
                    room.capacity = beds
                if room.capacity_estimated or not room.active or room.notes:
                    room.capacity_estimated, room.active, room.notes = False, True, ""
                    changes = changes or ["marked as confirmed"]
                if changes:
                    room.save()
                    report.updated.append(f"{label}: {', '.join(changes)}")
                else:
                    report.unchanged += 1
            keep.add(room.pk)
            people = [a.employee for a in room.assignments.select_related("employee")]
            wrong = [p.full_name for p in people if p.gender and p.gender != gender]
            if wrong:
                report.gender_conflicts.append({"room": label, "gender": gender, "people": wrong})
        for room in Room.objects.filter(building__kind=BuildingKind.COMPANY).exclude(pk__in=keep).select_related("building"):
            label = f"{room.building.name} {room.name}"
            count = room.assignments.count()
            if count:
                report.kept_with_people.append({"room": label, "people": count})
            else:
                report.removed.append(label)
                room.delete()
        for building in Building.objects.filter(kind=BuildingKind.COMPANY):
            if not building.rooms.exists():
                building.delete()
        if not dry_run:
            AuditService.log(event_type="accommodation.rooms_imported", module="accommodation", actor=actor, severity=AuditSeverity.INFO, title="Accommodation rooms updated from the tracker", description=f"{len(report.created)} added, {len(report.updated)} changed, {len(report.removed)} removed.", metadata={"created": report.created[:100], "removed": report.removed[:100], "kept_with_people": report.kept_with_people})
        else:
            transaction.set_rollback(True)
    return report


@dataclass
class ImportReport:
    dry_run: bool = True
    active_rows: int = 0
    placed_inside: int = 0
    outside_unplaced: int = 0
    no_accommodation: int = 0
    moved: int = 0
    unchanged: int = 0
    genders_set: int = 0
    unmatched_ids: list = field(default_factory=list)
    unreadable_rooms: list = field(default_factory=list)
    unknown_rooms: dict = field(default_factory=dict)
    gender_mismatches: list = field(default_factory=list)
    bed_conflicts: list = field(default_factory=list)
    over_capacity_rooms: list = field(default_factory=list)
    freed_beds_from_inactive: int = 0

    def as_dict(self):
        return self.__dict__.copy()


def _header_map(sheet_rows):
    return {str(name).strip().lower(): index for index, name in enumerate(sheet_rows[0]) if name is not None}


def _find_staff_sheet(workbook):
    for sheet in workbook.worksheets:
        headers = {str(c).strip().lower() for c in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)) if c is not None}
        if {"id", "room allocated", "employment status"} <= headers:
            return sheet
    raise ValueError("No sheet with the columns ID, EMPLOYMENT STATUS, Accommodation and ROOM ALLOCATED was found.")


def import_workbook(file, *, dry_run, actor=None):
    """Place staff from the staff spreadsheet: a room allocated = in that room and bed; Accommodation YES with
    no room = outside accommodation; NO = none. Rooms must already exist (tracker or Add Room): a room that
    is not one of ours leaves the person in company accommodation with no room, and is listed in the report."""
    sheet = _find_staff_sheet(load_workbook(file, data_only=True))
    cols = _header_map([next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))])
    missing = [name for name in ("id", "employment status", "accommodation", "room allocated") if name not in cols]
    if missing:
        raise ValueError(f"Missing column(s): {', '.join(missing)}.")
    report = ImportReport(dry_run=dry_run)
    rows = [row for row in sheet.iter_rows(min_row=2, values_only=True) if row[cols["id"]]]
    with transaction.atomic():
        for row in rows:
            active = str(row[cols["employment status"]]).strip().lower() == "active"
            label = row[cols["room allocated"]]
            has_label = bool(label and str(label).strip())
            if not active:
                if has_label:
                    report.freed_beds_from_inactive += 1
                    employee = Employee.objects.filter(employee_id=str(row[cols["id"]]).strip()).first()
                    if employee is not None and RoomAssignment.objects.filter(employee=employee).exists():
                        AccommodationService.unassign(employee, actor=actor)
                continue
            report.active_rows += 1
            employee_id = str(row[cols["id"]]).strip()
            employee = Employee.objects.filter(employee_id=employee_id).first()
            if employee is None:
                report.unmatched_ids.append(employee_id)
                continue
            if "gender" in cols:
                gender = gender_of(row[cols["gender"]])
                if gender and employee.gender != gender:
                    Employee.objects.filter(pk=employee.pk).update(gender=gender)
                    employee.gender = gender
                    report.genders_set += 1
            parsed = parse_room_label(label) if has_label else None
            if has_label and not parsed:
                report.unreadable_rooms.append({"employee_id": employee_id, "value": str(label).strip()})
            if parsed:
                building_name, room_name, bed = parsed
                room = find_room(building_name, room_name)
                if room is None:
                    key = f"{room_name}" if building_name == MAIN_HOSTEL else f"{building_name} {room_name}"
                    report.unknown_rooms[key] = report.unknown_rooms.get(key, 0) + 1
                    AccommodationService.mark_inside_without_room(employee, str(label).strip())
                    report.placed_inside += 1
                    continue
                current = RoomAssignment.objects.filter(employee=employee).select_related("room").first()
                if current and current.room_id == room.pk and current.bed_number == bed:
                    report.unchanged += 1
                    report.placed_inside += 1
                    sync_employee_flags(employee)
                    continue
                if room.gender and employee.gender and room.gender != employee.gender:
                    report.gender_mismatches.append({"employee_id": employee_id, "name": employee.full_name, "gender": employee.gender, "room": f"{room.building.name} {room.name}", "room_gender": room.gender})
                    AccommodationService.mark_inside_without_room(employee, str(label).strip())
                    report.placed_inside += 1
                    continue
                clash = RoomAssignment.objects.filter(room=room, bed_number=bed).exclude(employee=employee).first()
                if clash:
                    report.bed_conflicts.append({"employee_id": employee_id, "room": f"{room.building.name} {room.name}", "bed": bed, "already": clash.employee.employee_id})
                    bed = None
                if room.capacity is not None and (bed or 0) > room.capacity:
                    report.over_capacity_rooms.append(f"{room.building.name} {room.name}")
                if current:
                    report.moved += 1
                RoomAssignment.objects.update_or_create(employee=employee, defaults={"room": room, "bed_number": bed, "assigned_by": actor})
                sync_employee_flags(Employee.objects.get(pk=employee.pk))
                report.placed_inside += 1
            elif str(row[cols["accommodation"]]).strip().upper() == "YES":
                RoomAssignment.objects.filter(employee=employee).delete()
                AccommodationService.unassign(employee, actor=actor, outside=True)
                report.outside_unplaced += 1
            else:
                RoomAssignment.objects.filter(employee=employee).delete()
                AccommodationService.unassign(employee, actor=actor, outside=False)
                report.no_accommodation += 1
        report.over_capacity_rooms = sorted(set(report.over_capacity_rooms))
        if not dry_run:
            AuditService.log(event_type="accommodation.imported", module="accommodation", actor=actor, severity=AuditSeverity.INFO, title="Accommodation staff file imported", description=f"{report.placed_inside} placed in rooms, {report.outside_unplaced} outside, {report.no_accommodation} with none.", metadata={"placed_inside": report.placed_inside, "outside": report.outside_unplaced, "none": report.no_accommodation, "unknown_rooms": report.unknown_rooms, "unmatched_ids": report.unmatched_ids[:50]})
        else:
            transaction.set_rollback(True)
    return report


# ---- what the screen shows ---------------------------------------------------------------

def _natural(name):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", name)]


def build_overview():
    """Everything the accommodation screen needs in one go."""
    assignments = list(RoomAssignment.objects.select_related("employee", "employee__department", "room", "room__building"))
    by_room = defaultdict(list)
    for assignment in assignments:
        by_room[assignment.room_id].append(assignment)
    buildings = []
    total_capacity = total_occupied = 0
    genders = {g: {"rooms": 0, "beds": 0, "occupied": 0, "rooms_full": 0, "rooms_with_space": 0} for g in ("male", "female", "unset")}
    for building in Building.objects.prefetch_related("rooms"):
        room_rows = []
        for room in building.rooms.all():
            people = sorted(by_room.get(room.pk, []), key=lambda a: (a.bed_number is None, a.bed_number or 0, a.employee.employee_id))
            occupied = len(people)
            capacity = room.capacity
            if not room.active:
                state = "closed"
            elif capacity is None:
                state = "unknown"
            elif occupied > capacity:
                state = "over"
            elif occupied == capacity:
                state = "full"
            elif occupied == 0:
                state = "empty"
            else:
                state = "space"
            if room.active and building.kind == BuildingKind.COMPANY:
                total_occupied += occupied
                total_capacity += capacity or occupied
                bucket = genders[room.gender or "unset"]
                bucket["rooms"] += 1
                bucket["beds"] += capacity or occupied
                bucket["occupied"] += occupied
                bucket["rooms_full"] += state in ("full", "over")
                bucket["rooms_with_space"] += state in ("space", "empty")
            room_rows.append({
                "id": room.pk, "name": room.name, "gender": room.gender, "capacity": capacity, "capacity_estimated": room.capacity_estimated, "occupied": occupied,
                "free": None if capacity is None else max(capacity - occupied, 0), "state": state, "active": room.active, "notes": room.notes,
                "occupants": [{"employee": a.employee.pk, "employee_id": a.employee.employee_id, "name": a.employee.full_name, "gender": a.employee.gender, "bed": a.bed_number, "department": a.employee.department.name if a.employee.department else None} for a in people],
            })
        room_rows.sort(key=lambda r: _natural(r["name"]))
        buildings.append({"id": building.pk, "name": building.name, "kind": building.kind, "address": building.address, "rooms": room_rows, "occupied": sum(r["occupied"] for r in room_rows), "capacity": sum(r["capacity"] or 0 for r in room_rows if r["active"]), "full_rooms": sum(r["state"] == "full" for r in room_rows), "empty_rooms": sum(r["state"] == "empty" for r in room_rows), "rooms_with_space": sum(r["state"] == "space" for r in room_rows)})

    active = Employee.objects.filter(status="active")
    placed_ids = {a.employee_id for a in assignments}
    inside_placed = sum(1 for a in assignments if a.room.building.kind == BuildingKind.COMPANY and a.employee.status == "active")
    outside_placed = sum(1 for a in assignments if a.room.building.kind == BuildingKind.EXTERNAL and a.employee.status == "active")
    outside_people = active.filter(lives_in_external_accommodation=True).exclude(pk__in=placed_ids)
    inside_people = active.filter(lives_in_company_hostel=True).exclude(pk__in=placed_ids)
    active_count = active.count()
    by_gender = {}
    for key in ("male", "female"):
        b = genders[key]
        by_gender[key] = {**b, "free": max(b["beds"] - b["occupied"], 0),
                          "waiting_for_a_room": (inside_people | outside_people).filter(gender=key).count()}
    return {
        "summary": {
            "active_staff": active_count,
            "inside": inside_placed + inside_people.count(),
            "inside_with_room": inside_placed,
            "inside_without_room": inside_people.count(),
            "outside": outside_placed + outside_people.count(),
            "outside_placed": outside_placed,
            "outside_unplaced": outside_people.count(),
            "none": active_count - (inside_placed + inside_people.count() + outside_placed + outside_people.count()),
            "beds_total": total_capacity, "beds_occupied": total_occupied, "beds_free": max(total_capacity - total_occupied, 0),
            "rooms_full": sum(g["rooms_full"] for g in genders.values()),
            "rooms_with_space": sum(g["rooms_with_space"] for g in genders.values()),
            "rooms_empty": sum(b["empty_rooms"] for b in buildings if b["kind"] == "company"),
            "by_gender": by_gender,
        },
        "buildings": buildings,
    }
