import io

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase
from openpyxl import Workbook
from rest_framework.test import APIClient

from employees.models import Employee

from .models import Building, Room, RoomAssignment
from .services import AccommodationService, build_overview, import_room_setup, import_workbook, parse_room_label, split_room_name

HEADERS = ["ID", "NAME", "GENDER", "EMPLOYMENT STATUS", "Accommodation ", "ROOM ALLOCATED"]


def staff_workbook(rows):
    book = Workbook()
    sheet = book.active
    sheet.title = "Employee_DB"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    buffer.seek(0)
    return buffer


def tracker_workbook(male, female, name="tracker.xlsx"):
    """A workbook laid out like the accommodation tracker's Room Setup sheet: male table left, female table right."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Room Setup"
    sheet.append(["MALE ROOMS - edit this table to add/remove rooms", None, None, None, None, "FEMALE ROOMS - edit this table to add/remove rooms", None, None, None])
    sheet.append(["Room No", "Beds (capacity)", "Cumulative (auto)", "Start bed", None, "Room No", "Beds (capacity)", "Cumulative (auto)", "Start bed"])
    for index in range(max(len(male), len(female)) + 1):
        left = male[index] if index < len(male) else (None, None)
        right = female[index] if index < len(female) else (None, None)
        sheet.append([left[0], left[1], None, None, None, right[0], right[1], None, None])
    sheet.append(["Total male beds configured:", sum(b for _, b in male)])
    buffer = io.BytesIO()
    book.save(buffer)
    buffer.seek(0)
    buffer.name = name
    return buffer


def make_rooms():
    hostel = Building.objects.create(name="Main Hostel")
    return hostel, Room.objects.create(building=hostel, name="Room 301", gender="female", capacity=3), Room.objects.create(building=hostel, name="Room 001", gender="male", capacity=3)


class RoomNameTests(SimpleTestCase):
    def test_names_map_to_buildings(self):
        self.assertEqual(split_room_name("Room101"), ("Main Hostel", "Room 101"))
        self.assertEqual(split_room_name("Lodge 2 Rm001"), ("Lodge 2", "Rm001"))
        self.assertEqual(split_room_name("Security Rm 006"), ("Security Quarters", "Rm 006"))
        self.assertEqual(split_room_name("Admin G1"), ("Admin", "G1"))
        self.assertEqual(parse_room_label("Room 301 - Bed 1"), ("Main Hostel", "Room 301", 1))
        self.assertIsNone(parse_room_label("    "))


class RoomSetupImportTests(TestCase):
    male = [("Room 001", 14), ("Room101", 3), ("Lodge 2 Rm001", 6), ("Security Rm 005", 15)]
    female = [("Room 201", 8), ("Room 301", 2)]

    def test_rooms_are_created_with_gender_and_beds_and_the_total_row_is_ignored(self):
        report = import_room_setup(tracker_workbook(self.male, self.female), dry_run=False)
        self.assertEqual((report.rooms_in_file, report.male_rooms, report.male_beds, report.female_rooms, report.female_beds), (6, 4, 38, 2, 10))
        room = Room.objects.get(name="Room 101")  # "Room101" is normalised
        self.assertEqual((room.building.name, room.gender, room.capacity, room.capacity_estimated), ("Main Hostel", "male", 3, False))
        self.assertEqual(Room.objects.get(name="Room 201").gender, "female")
        self.assertEqual(Room.objects.get(building__name="Lodge 2", name="Rm001").capacity, 6)
        self.assertFalse(Room.objects.filter(name__istartswith="total").exists())

    def test_a_preview_changes_nothing(self):
        report = import_room_setup(tracker_workbook(self.male, self.female), dry_run=True)
        self.assertEqual(len(report.created), 6)
        self.assertFalse(Room.objects.exists())

    def test_rooms_not_in_the_file_are_removed_but_a_room_with_people_is_kept_and_reported(self):
        hostel = Building.objects.create(name="Main Hostel")
        Room.objects.create(building=hostel, name="Room 999", capacity=2)  # empty and not in the tracker
        busy = Room.objects.create(building=hostel, name="Room 998", capacity=2)
        person = Employee.objects.create(employee_id="1", first_name="A", last_name="B", status="active")
        AccommodationService.assign(person, busy, bed_number=1)
        report = import_room_setup(tracker_workbook(self.male, self.female), dry_run=False)
        self.assertIn("Main Hostel Room 999", report.removed)
        self.assertFalse(Room.objects.filter(name="Room 999").exists())
        self.assertEqual(report.kept_with_people, [{"room": "Main Hostel Room 998", "people": 1}])
        self.assertTrue(Room.objects.filter(name="Room 998").exists())

    def test_uploading_again_only_changes_what_differs(self):
        import_room_setup(tracker_workbook(self.male, self.female), dry_run=False)
        report = import_room_setup(tracker_workbook([("Room 001", 16)] + self.male[1:], self.female), dry_run=False)
        self.assertEqual((report.created, report.removed, report.unchanged), ([], [], 5))
        self.assertEqual(report.updated, ["Main Hostel Room 001: beds 14 to 16"])

    def test_a_file_without_a_room_setup_sheet_is_refused(self):
        book = Workbook()
        book.active.title = "Other"
        buffer = io.BytesIO()
        book.save(buffer)
        buffer.seek(0)
        buffer.name = "x.xlsx"
        with self.assertRaises(ValueError):
            import_room_setup(buffer, dry_run=True)


class StaffImportTests(TestCase):
    def setUp(self):
        import_room_setup(tracker_workbook([("Room 001", 3)], [("Room 301", 2)]), dry_run=False)
        for number, gender in (("000001", "female"), ("000002", "male"), ("000003", "male"), ("000004", "male"), ("000005", "male")):
            Employee.objects.create(employee_id=number, first_name="Staff", last_name=number, gender=gender)
        Employee.objects.filter(employee_id="000005").update(status="inactive")

    rows = [
        ["000001", "A", "Female", "Active", "YES", "Room 301 - Bed 1"],
        ["000002", "B", "Male", "Active", "YES", "Room 001 - Bed 2"],
        ["000003", "C", "Male", "Active", "YES", None],
        ["000004", "D", "Male", "Active", "NO", None],
        ["000005", "E", "Male", "Inactive", "YES", "Room 001 - Bed 1"],
        ["000999", "Nobody", "Male", "Active", "YES", "Room 001 - Bed 3"],
    ]

    def test_a_preview_changes_nothing(self):
        report = import_workbook(staff_workbook(self.rows), dry_run=True)
        self.assertEqual((report.placed_inside, report.outside_unplaced, report.no_accommodation), (2, 1, 1))
        self.assertEqual(report.unmatched_ids, ["000999"])
        self.assertFalse(RoomAssignment.objects.exists())

    def test_apply_places_people_in_the_existing_rooms_and_sets_the_flags(self):
        import_workbook(staff_workbook(self.rows), dry_run=False)
        one, two, three, four = (Employee.objects.get(employee_id=n) for n in ("000001", "000002", "000003", "000004"))
        self.assertEqual((one.room_assignment.room.name, one.room_assignment.bed_number), ("Room 301", 1))
        self.assertTrue(two.lives_in_company_hostel and not two.lives_in_external_accommodation)
        self.assertTrue(three.lives_in_external_accommodation and not three.lives_in_company_hostel)
        self.assertFalse(four.lives_in_company_hostel or four.lives_in_external_accommodation)
        self.assertFalse(RoomAssignment.objects.filter(employee__employee_id="000005").exists())

    def test_a_room_that_is_not_one_of_ours_is_never_created(self):
        report = import_workbook(staff_workbook([["000002", "B", "Male", "Active", "YES", "Room 110 - Bed 1"]]), dry_run=False)
        self.assertEqual(report.unknown_rooms, {"Room 110": 1})
        self.assertFalse(Room.objects.filter(name="Room 110").exists())
        person = Employee.objects.get(employee_id="000002")
        self.assertTrue(person.lives_in_company_hostel)  # still in company accommodation, room not recorded
        self.assertFalse(RoomAssignment.objects.filter(employee=person).exists())

    def test_a_man_in_a_womens_room_is_reported_not_placed(self):
        report = import_workbook(staff_workbook([["000002", "B", "Male", "Active", "YES", "Room 301 - Bed 1"]]), dry_run=False)
        self.assertEqual([m["employee_id"] for m in report.gender_mismatches], ["000002"])
        self.assertFalse(RoomAssignment.objects.exists())

    def test_gender_is_read_from_the_file(self):
        Employee.objects.filter(employee_id="000001").update(gender="")
        report = import_workbook(staff_workbook(self.rows[:1]), dry_run=False)
        self.assertEqual(report.genders_set, 1)
        self.assertEqual(Employee.objects.get(employee_id="000001").gender, "female")

    def test_importing_again_changes_nothing(self):
        import_workbook(staff_workbook(self.rows), dry_run=False)
        again = import_workbook(staff_workbook(self.rows), dry_run=False)
        self.assertEqual((again.unchanged, again.moved), (2, 0))


class AssignAndOverviewTests(TestCase):
    def setUp(self):
        self.hostel, self.women, self.men = make_rooms()
        self.a = Employee.objects.create(employee_id="1", first_name="Ada", last_name="A", gender="female")
        self.b = Employee.objects.create(employee_id="2", first_name="Bayo", last_name="B", gender="male")
        self.c = Employee.objects.create(employee_id="3", first_name="Chi", last_name="C", gender="male")

    def test_the_floors_are_kept_apart(self):
        with self.assertRaises(ValueError):
            AccommodationService.assign(self.b, self.women, bed_number=1)
        with self.assertRaises(ValueError):
            AccommodationService.assign(self.a, self.men, bed_number=1)
        AccommodationService.assign(self.a, self.women, bed_number=1)
        AccommodationService.assign(self.b, self.men, bed_number=1)

    def test_beds_cannot_be_double_booked_and_a_full_room_refuses_more(self):
        AccommodationService.assign(self.b, self.men, bed_number=1)
        with self.assertRaises(ValueError):
            AccommodationService.assign(self.c, self.men, bed_number=1)
        AccommodationService.assign(self.c, self.men, bed_number=2)
        Room.objects.filter(pk=self.men.pk).update(capacity=2)
        extra = Employee.objects.create(employee_id="9", first_name="X", last_name="X", gender="male")
        with self.assertRaises(ValueError):
            AccommodationService.assign(extra, self.men)

    def test_the_overview_shows_free_beds_for_the_male_and_female_floors(self):
        AccommodationService.assign(self.a, self.women, bed_number=1)
        AccommodationService.assign(self.b, self.men, bed_number=1)
        Employee.objects.filter(pk=self.c.pk).update(lives_in_external_accommodation=True)
        summary = build_overview()["summary"]
        self.assertEqual(summary["by_gender"]["female"], {"rooms": 1, "beds": 3, "occupied": 1, "rooms_full": 0, "rooms_with_space": 1, "free": 2, "waiting_for_a_room": 0})
        self.assertEqual(summary["by_gender"]["male"]["free"], 2)
        self.assertEqual(summary["by_gender"]["male"]["waiting_for_a_room"], 1)


class RoomApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        manager = get_user_model().objects.create_user("manager", password="pw")
        manager.user_permissions.add(Permission.objects.get(codename="manage_accommodation"))
        self.client.force_authenticate(manager)

    def test_add_a_room_then_remove_it(self):
        created = self.client.post("/api/accommodation/rooms/", {"building_name": "Lodge 3", "name": "Rm001", "gender": "male", "capacity": 4}, format="json")
        self.assertEqual(created.status_code, 201)
        room = Room.objects.get(name="Rm001")
        self.assertEqual((room.building.name, room.gender, room.capacity), ("Lodge 3", "male", 4))
        self.assertEqual(self.client.post("/api/accommodation/rooms/", {"building_name": "Lodge 3", "name": "rm001"}, format="json").status_code, 400)
        self.assertEqual(self.client.delete(f"/api/accommodation/rooms/{room.pk}/").status_code, 204)
        self.assertFalse(Room.objects.filter(name="Rm001").exists())

    def test_a_room_with_people_cannot_be_removed(self):
        _hostel, _women, men = make_rooms()
        person = Employee.objects.create(employee_id="2", first_name="Bayo", last_name="B", gender="male")
        AccommodationService.assign(person, men, bed_number=1)
        response = self.client.delete(f"/api/accommodation/rooms/{men.pk}/")
        self.assertEqual(response.status_code, 400)
        self.assertIn("still in", response.json()["detail"])
        self.assertTrue(Room.objects.filter(pk=men.pk).exists())

    def test_gender_and_beds_can_be_edited_but_not_to_strand_the_people_inside(self):
        _hostel, _women, men = make_rooms()
        person = Employee.objects.create(employee_id="2", first_name="Bayo", last_name="B", gender="male")
        AccommodationService.assign(person, men, bed_number=1)
        self.assertEqual(self.client.patch(f"/api/accommodation/rooms/{men.pk}/", {"gender": "female"}, format="json").status_code, 400)
        self.assertEqual(self.client.patch(f"/api/accommodation/rooms/{men.pk}/", {"capacity": 8}, format="json").status_code, 200)
        self.assertEqual(Room.objects.get(pk=men.pk).capacity, 8)

    def test_permissions_and_the_tracker_upload(self):
        outsider = APIClient()
        outsider.force_authenticate(get_user_model().objects.create_user("outsider", password="pw"))
        self.assertEqual(outsider.get("/api/accommodation/").status_code, 403)
        upload = SimpleUploadedFile("tracker.xlsx", tracker_workbook([("Room 001", 3)], [("Room 201", 2)]).read())
        preview = self.client.post("/api/accommodation/import-rooms/", {"file": upload}, format="multipart")
        self.assertEqual((preview.status_code, preview.json()["dry_run"], preview.json()["rooms_in_file"]), (200, True, 2))
        self.assertFalse(Room.objects.exists())
        staff = SimpleUploadedFile("staff.xlsx", staff_workbook([["000001", "A", "Female", "Active", "YES", None]]).read())
        self.assertEqual(self.client.post("/api/accommodation/import/", {"file": staff}, format="multipart").status_code, 200)
