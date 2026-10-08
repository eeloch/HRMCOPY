"""Load the contract staff's deferred fund balances and monthly percentages from the RAE Deferred Fund Contribution sheet.

Columns: Employee ID (the staff number, leading zeros dropped), Name, Current Balance (the deposit held), Contribution %
(0.15 or 15 both mean 15%). For each person: the staff record is found by staff number and the name is checked; if the
number does not exist the name is used when exactly one person has it. Anyone not already a contract employee is changed
to Contract. They are enrolled with the balance as an opening balance. People who have already left are recorded with
their balance but are not enrolled for further monthly deductions.

--contributions-from is the first payroll month the system deducts (the sheet's balance already includes the months
before it). Nothing is saved without --apply."""

import re
from datetime import date
from decimal import ROUND_HALF_UP, Decimal

from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.db.models import Q
from openpyxl import load_workbook

from audit.models import AuditSeverity
from audit.services import AuditService
from deferredfunds.models import DeferredFundAccount
from deferredfunds.services import DeferredFundService
from employees.models import Employee, EmploymentType


def name_key(value):
    return sorted(re.sub(r"[^a-z ]", "", (value or "").lower()).split())


def read_sheet(path):
    sheet = load_workbook(path, data_only=True).active
    rows = []
    for number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if values[0] is None and values[1] is None and values[2] is None:
            continue
        rows.append({"row": number, "id": values[0], "name": (values[1] or "").strip(), "balance": values[2], "percent": values[3]})
    return rows


def find_employee(row):
    """(employee, how) - by staff number with the name checked, else by a unique exact name."""
    try:
        number = f"{int(row['id']):06d}"
    except (TypeError, ValueError):
        number = ""
    employee = Employee.objects.filter(employee_id=number).first() if number else None
    if employee is not None and name_key(employee.full_name) == name_key(row["name"]):
        return employee, "staff number"
    same_name = [e for e in Employee.objects.all() if name_key(e.full_name) == name_key(row["name"])]
    if len(same_name) == 1:
        how = "name (the staff number in the sheet does not match)" if employee is None else "name (the sheet's staff number belongs to someone else)"
        return same_name[0], how
    return None, ""


class Command(BaseCommand):
    help = "Load deferred fund balances and percentages from the contract staff sheet (dry run unless --apply)."

    def add_arguments(self, parser):
        parser.add_argument("file")
        parser.add_argument("--contributions-from", default="2026-10-01", help="First payroll month the system deducts (YYYY-MM-DD).")
        parser.add_argument("--apply", action="store_true", help="Save. Without it nothing is changed.")
        parser.add_argument("--actor", default="admin", help="Username recorded as the person who loaded it.")

    def handle(self, *args, **options):
        rows = read_sheet(options["file"])
        start = date.fromisoformat(options["contributions_from"])
        actor = get_user_model().objects.filter(username=options["actor"]).first()
        if actor is None:
            raise CommandError(f"No user named {options['actor']}.")
        plan, problems, seen = [], [], set()
        for row in rows:
            employee, how = find_employee(row)
            if employee is None:
                problems.append(f"Row {row['row']}: {row['name']} (staff no. {row['id']}) - no unique match.")
                continue
            if employee.pk in seen:
                problems.append(f"Row {row['row']}: {row['name']} appears twice in the sheet.")
                continue
            seen.add(employee.pk)
            try:
                balance = Decimal(str(row["balance"])).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                percent = Decimal(str(row["percent"]))
                percent = (percent * 100 if percent <= 1 else percent).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            except Exception:
                problems.append(f"Row {row['row']}: {row['name']} - the balance or percentage is not a number.")
                continue
            if balance < 0 or not Decimal("0") < percent <= Decimal("100"):
                problems.append(f"Row {row['row']}: {row['name']} - balance {balance} / percentage {percent} is not valid.")
                continue
            if DeferredFundAccount.objects.filter(employee=employee).exists():
                problems.append(f"Row {row['row']}: {employee.full_name} already has a deferred fund account - left as it is.")
                continue
            plan.append((row, employee, how, balance, percent))
        retype = [item for item in plan if item[1].employment_type != EmploymentType.CONTRACT]
        self.stdout.write(f"{len(rows)} rows: {len(plan)} to load, {len(problems)} problem(s).")
        self.stdout.write(f"Employment type changed to Contract for {len(retype)} (now: " + ", ".join(f"{k} {v}" for k, v in sorted(self.count(item[1].employment_type for item in retype).items())) + ").")
        self.stdout.write(f"Exited staff recorded without monthly deductions: {sum(1 for item in plan if item[1].status != 'active')}. Total balance loaded: {sum(item[3] for item in plan):,.2f}.")
        for line in problems:
            self.stdout.write(self.style.WARNING(line))
        for row, employee, how, _balance, _percent in plan:
            if not how.startswith("staff number"):
                self.stdout.write(f"Matched {row['name']} to {employee.employee_id} by {how}.")
        if not options["apply"]:
            self.stdout.write(self.style.WARNING("DRY RUN - nothing was saved. Add --apply to save."))
            return
        with transaction.atomic():
            for row, employee, how, balance, percent in plan:
                if employee.employment_type != EmploymentType.CONTRACT:
                    before = employee.employment_type
                    employee.employment_type = EmploymentType.CONTRACT
                    employee.save(update_fields=["employment_type", "updated_at"])
                    AuditService.log(event_type="employees.employment_type_changed", module="employees", employee=employee, actor=actor, severity=AuditSeverity.WARNING,
                                     title="Employment type changed to Contract", description=f"{employee.full_name}: {before} to contract, because they are on the deferred fund contract staff sheet.",
                                     metadata={"from": before, "to": "contract"})
                account = DeferredFundService.enrol(employee=employee, percent=percent, opening_balance=balance, contributions_from=start, opening_note="Balance till Sep 2026 per the RAE Deferred Fund Contribution sheet", actor=actor)
                if employee.status != "active":
                    account.active = False  # has left: the balance is on record, no more monthly deductions
                    account.save(update_fields=["active", "updated_at"])
        self.stdout.write(self.style.SUCCESS(f"Loaded {len(plan)} deferred fund account(s)."))

    @staticmethod
    def count(values):
        counts = {}
        for value in values:
            counts[value] = counts.get(value, 0) + 1
        return counts
