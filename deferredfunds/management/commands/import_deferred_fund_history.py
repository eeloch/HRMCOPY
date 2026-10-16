"""Load the month-by-month deferred fund history (the "Deferred Fund Historical Data" workbook) into each person's ledger.

The workbook has two sheets: one row per person with what was contributed in every month, and a "Disbursement Record" of
what has been paid out since. Until now an account holds one line, "Opening balance", for the whole of the past. This
replaces that line with the real history - a contribution line for each month and a line for what was paid out - so the
statement shows how the balance was built up. THE BALANCE DOES NOT CHANGE: an account is only converted when the history
adds up to exactly the opening balance it already holds (total saved less total disbursed). Anyone whose figures differ is
listed and left untouched for a person to decide. Accounts that already carry history are skipped, so it is safe to run twice.

People are found by staff number with the name checked (the staff number appears twice in the sheet for two different
people, so the disbursement sheet is matched on number AND name). The month columns are dated as the sheet shows them; a
header written as text ("January 26") is dated the 26th. The disbursement sheet has no dates: --disbursement-date says what
date to give those lines. The date also gets saving_since set to the first month saved, when it is still blank, because
the 18-month rule counts from it.

Nothing is saved without --apply."""

import re
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal

from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from audit.models import AuditSeverity
from audit.services import AuditService
from deferredfunds.management.commands.import_deferred_fund_sheet import find_employee, name_key
from deferredfunds.models import DeferredFundAccount, DeferredFundEntry, EntryType
from employees.models import Employee

TWO = Decimal("0.01")
HISTORY_TAG = "(history)"


def money(value):
    return Decimal(str(value)).quantize(TWO, rounding=ROUND_HALF_UP)


def month_date(header):
    """The date a month column stands for, or None when the header is not a month."""
    if isinstance(header, datetime):
        return header.date()
    if isinstance(header, date):
        return header
    text = str(header or "").strip()
    for pattern in ("%B %y", "%B %Y", "%b %y", "%b %Y"):
        try:
            parsed = datetime.strptime(text, pattern)
            return date(parsed.year, parsed.month, 26)
        except ValueError:
            continue
    return None


def read_history(path):
    """(people, disbursements). people: one dict per row with the months; disbursements: dicts with number, name, amount."""
    workbook = load_workbook(path, data_only=True)
    record = next((ws for ws in workbook.worksheets if "disburs" not in ws.title.lower()), None)
    paid = next((ws for ws in workbook.worksheets if "disburs" in ws.title.lower()), None)
    if record is None:
        raise CommandError("The workbook has no history sheet.")
    rows = list(record.iter_rows(values_only=True))
    header = [str(h).strip().lower() if isinstance(h, str) else h for h in rows[0]]

    def column(*names):
        for position, value in enumerate(header):
            if isinstance(value, str) and value in names:
                return position
        raise CommandError(f"The history sheet has no column called {names[0]!r}.")

    number, name, percent, status = column("employee number", "employee id"), column("name"), column("fund contribution"), column("status")
    months = [(position, month_date(value)) for position, value in enumerate(rows[0]) if month_date(value) is not None]
    if not months:
        raise CommandError("The history sheet has no month columns.")
    people = []
    for line, values in enumerate(rows[1:], start=2):
        if values[number] is None and values[name] is None:
            continue
        people.append({
            "row": line, "id": values[number], "name": (values[name] or "").strip(), "percent": values[percent], "status": values[status],
            "months": [(day, values[position]) for position, day in months],
        })
    disbursements = []
    if paid is not None:
        for line, values in enumerate(paid.iter_rows(min_row=2, values_only=True), start=2):
            if all(v is None for v in values[:3]):
                continue
            disbursements.append({"row": line, "id": values[0], "name": (values[1] or "").strip(), "amount": values[2]})
    return people, disbursements


class Command(BaseCommand):
    help = "Replace each account's single opening-balance line with the monthly history from the workbook (dry run unless --apply)."

    def add_arguments(self, parser):
        parser.add_argument("file")
        parser.add_argument("--disbursement-date", help="Date (YYYY-MM-DD) to give the disbursement lines; the sheet has none. Needed with --apply.")
        parser.add_argument("--create-missing", action="store_true", help="Also create an inactive account (no monthly deductions) for people in the sheet who have none, so their history is on record.")
        parser.add_argument("--apply", action="store_true", help="Save. Without it nothing is changed.")
        parser.add_argument("--actor", default="admin", help="Username recorded as the person who loaded it.")

    def handle(self, *args, **options):
        people, disbursements = read_history(options["file"])
        actor = get_user_model().objects.filter(username=options["actor"]).first()
        if actor is None:
            raise CommandError(f"No user named {options['actor']}.")
        disbursed_on = None
        if options["disbursement_date"]:
            try:
                disbursed_on = date.fromisoformat(options["disbursement_date"])
            except ValueError:
                raise CommandError("--disbursement-date must be YYYY-MM-DD.")

        problems, plan, skipped = [], [], []
        seen = set()
        paid_out = {}
        for item in disbursements:
            try:
                amount = money(item["amount"])
            except Exception:
                problems.append(f"Disbursement row {item['row']}: {item['name']} - the amount is not a number.")
                continue
            if amount <= 0:
                problems.append(f"Disbursement row {item['row']}: {item['name']} - the amount must be more than zero.")
                continue
            paid_out.setdefault((str(item["id"]), tuple(name_key(item["name"]))), []).append((item, amount))

        for person in people:
            employee, how = find_employee(person)
            if employee is None:
                problems.append(f"Row {person['row']}: {person['name']} (staff no. {person['id']}) - no unique match.")
                continue
            if employee.pk in seen:
                problems.append(f"Row {person['row']}: {person['name']} appears twice in the sheet.")
                continue
            seen.add(employee.pk)
            contributions, bad = [], False
            for day, value in person["months"]:
                if value in (None, ""):
                    continue
                try:
                    amount = money(value)
                except Exception:
                    problems.append(f"Row {person['row']}: {person['name']} - {day:%b %Y} is not a number.")
                    bad = True
                    break
                if amount < 0:
                    problems.append(f"Row {person['row']}: {person['name']} - {day:%b %Y} is negative.")
                    bad = True
                    break
                if amount > 0:
                    contributions.append((day, amount))
            if bad:
                continue
            key = (str(person["id"]), tuple(name_key(person["name"])))
            payments = paid_out.pop(key, [])
            gross = sum((a for _, a in contributions), Decimal("0.00"))
            disbursed = sum((a for _, a in payments), Decimal("0.00"))
            net = gross - disbursed
            label = f"Row {person['row']}: {employee.employee_id} {employee.full_name}"
            if not contributions and not payments:
                skipped.append(f"{label} - nothing was saved in any month.")
                continue
            if net < 0:
                skipped.append(f"{label} - paid out {disbursed:,.2f} but only {gross:,.2f} was saved ({-net:,.2f} more than the history holds). Left for a person to decide.")
                continue
            account = DeferredFundAccount.objects.filter(employee=employee).first()
            if account is None:
                if options["create_missing"]:
                    plan.append({"person": person, "employee": employee, "account": None, "contributions": contributions, "payments": payments, "gross": gross, "disbursed": disbursed, "net": net, "label": label})
                else:
                    skipped.append(f"{label} - no deferred fund account (saved {gross:,.2f}, paid out {disbursed:,.2f}). Add --create-missing to record the history on an inactive account.")
                continue
            entries = list(account.entries.all())
            if any(e.note.endswith(HISTORY_TAG) for e in entries):
                skipped.append(f"{label} - the history is already loaded.")
                continue
            openings = [e for e in entries if e.entry_type == EntryType.OPENING]
            if len(openings) != 1:
                skipped.append(f"{label} - {'no opening balance line' if not openings else 'more than one opening balance line'} to replace (history nets {net:,.2f}). Left for a person to decide.")
                continue
            if openings[0].amount != net:
                skipped.append(f"{label} - the opening balance is {openings[0].amount:,.2f} but the history nets {net:,.2f} (saved {gross:,.2f} less paid out {disbursed:,.2f}); difference {net - openings[0].amount:+,.2f}. Left for a person to decide.")
                continue
            plan.append({"person": person, "employee": employee, "account": account, "opening": openings[0], "contributions": contributions, "payments": payments, "gross": gross, "disbursed": disbursed, "net": net, "label": label})

        for (number, _), items in paid_out.items():
            for item, amount in items:
                problems.append(f"Disbursement row {item['row']}: {item['name']} (staff no. {number}) {amount:,.2f} - no matching person on the history sheet.")

        needs_date = [p for p in plan if p["payments"]]
        self.stdout.write(f"{len(people)} people on the history sheet, {len(disbursements)} disbursement line(s).")
        self.stdout.write(f"To load: {len(plan)} ({sum(1 for p in plan if p['account'] is None)} on new inactive accounts); saved {sum((p['gross'] for p in plan), Decimal('0')):,.2f}, paid out {sum((p['disbursed'] for p in plan), Decimal('0')):,.2f}, held {sum((p['net'] for p in plan), Decimal('0')):,.2f}.")
        self.stdout.write(f"Left untouched: {len(skipped)}. Problems: {len(problems)}.")
        for line in skipped:
            self.stdout.write(self.style.WARNING(line))
        for line in problems:
            self.stdout.write(self.style.WARNING(line))
        if needs_date and disbursed_on is None:
            self.stdout.write(self.style.WARNING(f"{len(needs_date)} account(s) have disbursements and the sheet gives no date: give --disbursement-date YYYY-MM-DD before --apply."))
        if not options["apply"]:
            self.stdout.write(self.style.WARNING("DRY RUN - nothing was saved. Add --apply to save."))
            return
        if needs_date and disbursed_on is None:
            raise CommandError("Say what date the disbursements were made with --disbursement-date YYYY-MM-DD.")

        with transaction.atomic():
            for item in plan:
                self.load(item, actor, disbursed_on)
        self.stdout.write(self.style.SUCCESS(f"Loaded the history for {len(plan)} account(s)."))

    @staticmethod
    def load(item, actor, disbursed_on):
        employee, account = item["employee"], item["account"]
        first_month = item["contributions"][0][0] if item["contributions"] else None
        if account is None:
            percent = Decimal(str(item["person"]["percent"] or 0))
            percent = money(percent * 100 if percent <= 1 else percent)  # 0.15 and 15 both mean 15%
            account = DeferredFundAccount.objects.create(employee=employee, percent=percent if percent > 0 else Decimal("0.01"), active=False, saving_since=first_month, created_by=actor)
        else:
            item["opening"].delete()
            if account.saving_since is None and first_month is not None:
                account.saving_since = first_month
                account.save(update_fields=["saving_since", "updated_at"])
        for day, amount in item["contributions"]:
            DeferredFundEntry.objects.create(account=account, entry_type=EntryType.CONTRIBUTION, amount=amount, entry_date=day, note=f"{day:%b %Y} contribution {HISTORY_TAG}", created_by=actor)
        total = len(item["payments"])
        for position, (_, amount) in enumerate(item["payments"], start=1):
            final = item["net"] == 0 and position == total
            DeferredFundEntry.objects.create(
                account=account, entry_type=EntryType.RELEASE if final else EntryType.WITHDRAWAL, amount=-amount, entry_date=disbursed_on,
                note=f"Paid out before this system {HISTORY_TAG}", created_by=actor,
            )
        AuditService.log(
            event_type="deferred_fund.history_loaded", module="deferred_funds", employee=employee, actor=actor, object=account, severity=AuditSeverity.INFO,
            title="Deferred fund history loaded",
            description=f"{employee.full_name}: {len(item['contributions'])} monthly contribution(s) totalling {item['gross']:,.2f} and {item['disbursed']:,.2f} paid out now make up the balance of {item['net']:,.2f}.",
            metadata={"account_id": account.pk, "saved": str(item["gross"]), "paid_out": str(item["disbursed"]), "balance": str(item["net"])},
        )
