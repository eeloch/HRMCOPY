import io
import zipfile
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import SimpleTestCase, TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from audit.models import AuditEvent
from employees.banking import normalize_bank_code, to_account_number, to_bank_code
from employees.models import Employee
from payroll.models import EmployeePayroll, EmployeePayrollStatus, PayrollPeriod


class BankDetailNormalisationTests(SimpleTestCase):
    def test_bank_codes_get_their_leading_zeros_back(self):
        self.assertEqual([normalize_bank_code(c) for c in ("14", "4", "1", "90267", "100004", " 000015 ")], ["000014", "000004", "000001", "090267", "100004", "000015"])

    def test_a_bank_code_that_is_not_all_digits_is_left_alone(self):
        self.assertEqual(normalize_bank_code("GTB"), "GTB")
        self.assertEqual(to_bank_code("GTB"), (None, "Bank code 'GTB' is not a 6-digit code"))

    def test_bank_code_problems(self):
        self.assertEqual(to_bank_code(""), (None, "No bank code"))
        self.assertEqual(to_bank_code("1234567")[0], None)
        self.assertEqual(to_bank_code("14"), ("000014", None))

    def test_a_ten_digit_account_is_used_as_is(self):
        self.assertEqual(to_account_number("2252037955"), ("2252037955", False, None))

    def test_spaces_including_non_breaking_ones_and_hyphens_are_dropped(self):
        self.assertEqual(to_account_number("913\xa0221\xa07704"), ("9132217704", False, None))
        self.assertEqual(to_account_number("225-203-7955"), ("2252037955", False, None))

    def test_eight_and_nine_digit_accounts_are_padded_but_flagged(self):
        self.assertEqual(to_account_number("68573701"), ("0068573701", True, None))
        self.assertEqual(to_account_number("147400645"), ("0147400645", True, None))

    def test_accounts_that_cannot_be_a_real_number_are_a_problem_not_a_guess(self):
        self.assertEqual(to_account_number("")[2], "No account number")
        self.assertIn("11 digits", to_account_number("11773623135")[2])
        self.assertIn("only 7 digits", to_account_number("1234567")[2])
        self.assertIn("non-digit", to_account_number("12345ABC90")[2])


class BankUploadTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.manager = get_user_model().objects.create_user("payroll-manager", password="pw")
        self.manager.user_permissions.add(Permission.objects.get(codename="manage_payroll"))
        self.period = PayrollPeriod.objects.create(year=2026, month=9, status="approved")
        self.client.force_authenticate(self.manager)

    def person(self, staff_id, net, account="2252037955", code="000015", bank="Zenith Bank Plc", first="Ada"):
        employee = Employee.objects.create(employee_id=staff_id, first_name=first, last_name=staff_id, bank_name=bank, account_number=account, bank_code=code)
        EmployeePayroll.objects.create(payroll_period=self.period, employee=employee, basic_salary=net, gross_earnings=net, net_pay=net, status=EmployeePayrollStatus.APPROVED)
        return employee

    def sheet(self, response):
        return load_workbook(io.BytesIO(response.content)).active

    def test_the_workbook_matches_the_format_hr_uploads_today(self):
        self.person("000001", Decimal("340000.00"), account="1539197070", code="14")
        self.person("000002", Decimal("28962.33"), account="21254638", code="000004")

        response = self.client.get(f"/api/payroll/periods/{self.period.pk}/bank-upload/")

        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn("September 2026 - Bank Upload.xlsx", response["Content-Disposition"])
        sheet = self.sheet(response)
        self.assertEqual(sheet.title, "September 2026")
        self.assertEqual([c.value for c in sheet[1]], ["Account_Number", "Amount", "Bank_Codes", "Narration"])
        self.assertEqual(sheet.max_row, 3)  # header + 2 people, no totals row
        first, second = sheet[2], sheet[3]
        self.assertEqual([c.value for c in first], [1539197070, 340000, "000014", "Hello"])
        self.assertEqual([c.value for c in second], [21254638, 28962.33, "000004", "Hello"])
        self.assertEqual([c.number_format for c in first], ["0000000000", "#,##0.00", "@", "General"])
        self.assertIsInstance(first[0].value, int)
        self.assertIsInstance(first[2].value, str)

    def test_every_row_carries_the_narration(self):
        for n in range(3):
            self.person(f"00000{n}", Decimal("50000.00"))
        rows = list(self.sheet(self.client.get(f"/api/payroll/periods/{self.period.pk}/bank-upload/")).iter_rows(min_row=2, values_only=True))
        self.assertEqual({row[3] for row in rows}, {"Hello"})

    def test_people_who_cannot_be_paid_are_left_out_and_explained(self):
        self.person("000001", Decimal("100000.00"))
        self.person("000002", Decimal("100000.00"), account="")
        self.person("000003", Decimal("100000.00"), code="")
        self.person("000004", Decimal("0.00"))

        preview = self.client.get(f"/api/payroll/periods/{self.period.pk}/bank-upload/preview/").json()

        self.assertEqual((preview["ready_count"], preview["ready_total"], preview["narration"]), (1, "100000.00", "Hello"))
        reasons = {i["employee_id"]: (i["reason"], i["fixable"]) for i in preview["issues"]}
        self.assertEqual(reasons["000002"], ("No account number", True))
        self.assertEqual(reasons["000003"], ("No bank code", True))
        self.assertEqual(reasons["000004"][1], False)
        self.assertIn("nothing to pay", reasons["000004"][0])
        sheet = self.sheet(self.client.get(f"/api/payroll/periods/{self.period.pk}/bank-upload/"))
        self.assertEqual(sheet.max_row, 2)

    def test_a_padded_account_is_included_but_listed_for_a_double_check(self):
        self.person("000001", Decimal("100000.00"), account="68573701")

        preview = self.client.get(f"/api/payroll/periods/{self.period.pk}/bank-upload/preview/").json()

        self.assertEqual(preview["ready_count"], 1)
        self.assertEqual(preview["padded"], [{"employee_id": "000001", "employee_name": "Ada 000001", "account_number": "0068573701"}])

    def test_it_is_only_available_once_the_payroll_is_approved(self):
        self.period.status = "draft"
        self.period.save(update_fields=["status"])
        self.person("000001", Decimal("100000.00"))
        for url in ("bank-upload/", "bank-upload/preview/"):
            self.assertEqual(self.client.get(f"/api/payroll/periods/{self.period.pk}/{url}").status_code, 400)

    def test_it_needs_the_manage_payroll_permission(self):
        self.person("000001", Decimal("100000.00"))
        viewer = get_user_model().objects.create_user("payroll-viewer", password="pw")
        viewer.user_permissions.add(Permission.objects.get(codename="view_payroll"))
        self.client.force_authenticate(viewer)
        for url in ("bank-upload/", "bank-upload/preview/"):
            self.assertEqual(self.client.get(f"/api/payroll/periods/{self.period.pk}/{url}").status_code, 403)

    def test_nothing_to_export_is_a_clear_error_not_an_empty_file(self):
        self.person("000001", Decimal("100000.00"), account="")
        response = self.client.get(f"/api/payroll/periods/{self.period.pk}/bank-upload/")
        self.assertEqual(response.status_code, 400)

    def test_a_batch_size_splits_the_file_into_a_zip_of_batches(self):
        for n in range(5):
            self.person(f"00000{n}", Decimal("50000.00"), account=f"22520379{n}5")

        response = self.client.get(f"/api/payroll/periods/{self.period.pk}/bank-upload/?batch_size=2")

        self.assertEqual(response["Content-Type"], "application/zip")
        bundle = zipfile.ZipFile(io.BytesIO(response.content))
        self.assertEqual(bundle.namelist(), ["Batch 1.xlsx", "Batch 2.xlsx", "Batch 3.xlsx"])
        sizes = [load_workbook(io.BytesIO(bundle.read(name))).active.max_row - 1 for name in bundle.namelist()]
        self.assertEqual(sizes, [2, 2, 1])

    def test_a_batch_size_bigger_than_the_list_still_gives_a_single_workbook(self):
        self.person("000001", Decimal("100000.00"))
        response = self.client.get(f"/api/payroll/periods/{self.period.pk}/bank-upload/?batch_size=100")
        self.assertIn("spreadsheetml", response["Content-Type"])

    def test_a_bad_batch_size_is_rejected(self):
        self.person("000001", Decimal("100000.00"))
        for value in ("abc", "0", "999999"):
            self.assertEqual(self.client.get(f"/api/payroll/periods/{self.period.pk}/bank-upload/?batch_size={value}").status_code, 400)

    def test_creating_the_file_is_audit_logged(self):
        self.person("000001", Decimal("100000.00"))
        self.person("000002", Decimal("100000.00"), account="")

        self.client.get(f"/api/payroll/periods/{self.period.pk}/bank-upload/")

        event = AuditEvent.objects.get(event_type="payroll.bank_upload_exported")
        self.assertEqual((event.metadata["payments"], event.metadata["total"], event.metadata["left_out"]), (1, "100000.00", 1))

    def test_saving_an_employee_repairs_a_bank_code_that_lost_its_zeros(self):
        employee = Employee.objects.create(employee_id="000009", first_name="A", last_name="B", bank_code="14")
        employee.refresh_from_db()
        self.assertEqual(employee.bank_code, "000014")


class ApprovingAPeriodApprovesItsEmployeeRecordsTests(TestCase):
    def test_transition_to_approved_moves_draft_employee_records_to_approved(self):
        manager = get_user_model().objects.create_user("approver", password="pw")
        manager.user_permissions.add(Permission.objects.get(codename="manage_payroll"))
        period = PayrollPeriod.objects.create(year=2026, month=9, status="review")
        employee = Employee.objects.create(employee_id="000001", first_name="Ada", last_name="Okafor")
        payroll = EmployeePayroll.objects.create(payroll_period=period, employee=employee, basic_salary=1, gross_earnings=1, net_pay=1)
        client = APIClient(); client.force_authenticate(manager)

        response = client.post(f"/api/payroll/periods/{period.pk}/transition/", {"status": "approved"}, format="json")

        self.assertEqual(response.status_code, 200, response.content)
        payroll.refresh_from_db()
        self.assertEqual(payroll.status, EmployeePayrollStatus.APPROVED)
