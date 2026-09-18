"""The bank upload spreadsheet for an approved payroll period.

Mirrors the sheet HR already uploads to the bank ("Batch 1.xlsx"): one header row
(Account_Number, Amount, Bank_Codes, Narration) then one row per person, no totals
row. The account is a number shown with 10 digits, the amount a number shown as
#,##0.00, the bank code text, and the narration is always the same word.

Only people who can actually be paid go in the file. Everyone else is returned as
an issue with the reason, so nobody is silently left out of the salary run.
"""

import io
import zipfile
from dataclasses import dataclass
from decimal import Decimal

from openpyxl import Workbook

from employees.banking import to_account_number, to_bank_code

BANK_UPLOAD_NARRATION = "Hello"
HEADERS = ("Account_Number", "Amount", "Bank_Codes", "Narration")
FINAL_STATUSES = {"approved", "paid", "closed"}


@dataclass
class BankUploadRow:
    employee_id: str
    employee_name: str
    account_number: str
    amount: Decimal
    bank_code: str
    padded: bool


@dataclass
class BankUploadIssue:
    employee_id: str
    employee_name: str
    net_pay: Decimal
    bank_name: str
    account_number: str
    bank_code: str
    reason: str
    fixable: bool


def prepare_bank_upload(period):
    """Split a period's payroll into (rows ready to upload, issues) in staff-number order."""
    rows, issues = [], []
    payrolls = period.employee_payrolls.select_related("employee").order_by("employee__employee_id")
    for payroll in payrolls:
        employee = payroll.employee
        common = dict(
            employee_id=employee.employee_id, employee_name=employee.full_name, net_pay=payroll.net_pay,
            bank_name=employee.bank_name, account_number=employee.account_number, bank_code=employee.bank_code,
        )
        if payroll.net_pay <= 0:
            issues.append(BankUploadIssue(**common, reason=f"Net pay is {payroll.net_pay:,.2f} - nothing to pay", fixable=False))
            continue
        account, padded, account_problem = to_account_number(employee.account_number)
        code, code_problem = to_bank_code(employee.bank_code)
        problems = [problem for problem in (account_problem, code_problem) if problem]
        if problems:
            issues.append(BankUploadIssue(**common, reason="; ".join(problems), fixable=True))
            continue
        rows.append(BankUploadRow(employee.employee_id, employee.full_name, account, payroll.net_pay.quantize(Decimal("0.01")), code, padded))
    return rows, issues


def build_workbook(rows, sheet_title):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title[:31]
    sheet.append(list(HEADERS))
    for index, row in enumerate(rows, start=2):
        sheet.append([int(row.account_number), float(row.amount), row.bank_code, BANK_UPLOAD_NARRATION])
        sheet.cell(index, 1).number_format = "0000000000"
        sheet.cell(index, 2).number_format = "#,##0.00"
        sheet.cell(index, 3).number_format = "@"
    sheet.column_dimensions["A"].width = 14.5
    sheet.column_dimensions["B"].width = 10.8
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def bank_upload_download(period, rows, batch_size=None):
    """Return (filename, content_type, bytes): one .xlsx, or a .zip of "Batch N.xlsx" files
    when a batch size is given and there are more rows than fit in one."""
    title = period.display_name
    if not batch_size or len(rows) <= batch_size:
        return (
            f"{title} - Bank Upload.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            build_workbook(rows, title),
        )
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as bundle:
        for number, start in enumerate(range(0, len(rows), batch_size), start=1):
            bundle.writestr(f"Batch {number}.xlsx", build_workbook(rows[start:start + batch_size], f"{title} Batch {number}"))
    return f"{title} - Bank Upload.zip", "application/zip", archive.getvalue()
